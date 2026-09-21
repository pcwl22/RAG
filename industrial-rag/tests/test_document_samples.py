"""Real sample document parsing tests."""
from app.parser.document_parser import parse_document


def test_parse_real_docx_sample(tmp_path):
    from docx import Document

    path = tmp_path / "sample.docx"
    document = Document()
    document.add_paragraph("First paragraph from docx.")
    document.add_paragraph("Second paragraph from docx.")
    document.save(path)

    parsed = parse_document(str(path))

    assert "First paragraph from docx." in parsed
    assert "Second paragraph from docx." in parsed


def test_parse_real_docx_sample_keeps_tables_in_document_order(tmp_path):
    from docx import Document

    path = tmp_path / "table.docx"
    document = Document()
    document.add_paragraph("Before table")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "name"
    table.cell(0, 1).text = "status"
    table.cell(1, 0).text = "pump"
    table.cell(1, 1).text = "ok"
    document.add_paragraph("After table")
    document.save(path)

    parsed = parse_document(str(path))

    assert parsed.index("Before table") < parsed.index("| name | status |") < parsed.index("After table")
    assert "| pump | ok |" in parsed


def test_parse_real_xlsx_sample(tmp_path):
    import pandas as pd

    path = tmp_path / "sample.xlsx"
    frame = pd.DataFrame(
        [
            {"name": "pump", "status": "ok"},
            {"name": "valve", "status": "check"},
        ]
    )
    frame.to_excel(path, index=False, sheet_name="assets")

    parsed = parse_document(str(path))

    assert "## assets" in parsed
    assert "name" in parsed
    assert "pump" in parsed
    assert "valve" in parsed


def test_parse_real_pdf_sample(tmp_path):
    import fitz

    path = tmp_path / "sample.pdf"
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), "PDF sample text for parser.")
    document.save(path)
    document.close()

    parsed = parse_document(str(path))

    assert "PDF sample text for parser." in parsed
